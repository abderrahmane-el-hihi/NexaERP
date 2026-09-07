from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "report-source.md"
OUTPUT = ROOT / "NexaERP_Morocco_Market_and_Production_Readiness_Report.docx"

NAVY = "17324D"
BLUE = "315F87"
PALE = "EAF1F7"
GRAY = "D9E1E8"
TEXT = "1F2933"
MUTED = "52616B"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    elem = OxmlElement("w:tblHeader")
    elem.set(qn("w:val"), "true")
    tr_pr.append(elem)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, font_size=10.2, color=TEXT):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        clean = part[2:-2] if bold else part
        run = paragraph.add_run(clean)
        run.bold = bold
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor.from_string(color)


def add_decision_table(doc):
    title = doc.add_paragraph()
    title.style = "Heading 1"
    title.add_run("Decision summary")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Decision", "Recommendation", "Evidence required"]
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        set_cell_border(cell, "D9E1E8")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("Commercial launch", "Run paid design-partner pilots; do not make legal-compliance claims.", "Signed pilots, cabinet participation, activation and support metrics."),
        ("Compliance launch", "Block until an authoritative DGI or approved-operator route is live and tested.", "Written requirements, credentials, contract tests, legal review and beta evidence."),
        ("First segment", "Commerce/distribution SMEs in Casablanca and Rabat, sold with their fiduciaire.", "Twenty structured interviews and five paid pilots."),
        ("Product sequencing", "Invoice control, AR, cash, stock and migration before payroll, manufacturing or generic AI.", "Repeatable end-to-end workflow and accountant sign-off."),
    ]
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cell = cells[j]
            set_cell_shading(cell, "F6F9FC" if idx % 2 else "FFFFFF")
            set_cell_margins(cell, top=110, bottom=110)
            set_cell_border(cell, "D9E1E8")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, 8.8)
    doc.add_paragraph()


def add_plan_table(doc):
    title = doc.add_paragraph()
    title.style = "Heading 1"
    title.add_run("Ninety day execution plan")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Period", "Outcome", "Non negotiable work"]
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        set_cell_border(cell, "D9E1E8")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
    set_repeat_table_header(table.rows[0])
    rows = [
        ("Days 1 to 30", "Regulatory path and customer problem validated", "Written DGI/operator/trust-provider answers; 20 interviews; 3 pilot commitments; truthful demo and landing page."),
        ("Days 31 to 60", "One repeatable pilot workflow", "Migration templates; invoice to payment flow; cabinet export; weekly support measurement; fix browser/onboarding defects."),
        ("Days 61 to 90", "Go/no-go evidence", "Five paid pilots or a documented pivot; real integration test route; production operations and privacy workstream."),
    ]
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cell = cells[j]
            set_cell_shading(cell, "F6F9FC" if idx % 2 else "FFFFFF")
            set_cell_margins(cell, top=110, bottom=110)
            set_cell_border(cell, "D9E1E8")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            add_inline(cell.paragraphs[0], value, 8.8)
    doc.add_paragraph()


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.13

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(9)

    for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("NexaERP Morocco Market and Production Readiness Report | 7 September 2026")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_cover(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("NexaERP Morocco Market and Production Readiness Report")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(20)
    add_inline(sub, "Go to market strategy, product scope and launch gates for Moroccan small and medium enterprises", 13, MUTED)

    meta = doc.add_table(rows=3, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.style = "Table Grid"
    values = [("Audience", "Founder and leadership team"), ("Date", "7 September 2026"), ("Primary decision", "What to sell now, what to build next, and when to permit general availability")]
    for row, (left, right) in zip(meta.rows, values):
        for cell in row.cells:
            set_cell_border(cell, "D9E1E8")
            set_cell_margins(cell, top=110, bottom=110)
        set_cell_shading(row.cells[0], PALE)
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = row.cells[0].paragraphs[0].add_run(left)
        r.bold = True
        r.font.size = Pt(9)
        add_inline(row.cells[1].paragraphs[0], right, 9)
    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(8)
    add_inline(intro, "Bottom line: NexaERP can run a focused paid-pilot motion now, but compliance general availability must remain blocked until a documented legal and technical route is tested in the real ecosystem.", 12)
    doc.add_page_break()


def build_document():
    source = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure(doc)
    add_cover(doc)
    add_decision_table(doc)

    source_section = None
    inserted_plan = False
    for raw in source:
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("**Audience:") or line.startswith("**Date:") or line.startswith("**Decision:"):
            continue
        if line.startswith("## "):
            heading = line[3:]
            if heading == "Go-to-market plan" and not inserted_plan:
                add_plan_table(doc)
                inserted_plan = True
            p = doc.add_paragraph(style="Heading 1")
            p.add_run(heading)
            set_keep_with_next(p)
            source_section = heading
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            p.add_run(line[4:])
            set_keep_with_next(p)
            continue
        if line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line[2:], 9.7)
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, re.sub(r"^\d+\. ", "", line), 9.7)
            continue
        if line.startswith("###"):
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            p.add_run(line[5:])
            set_keep_with_next(p)
            continue
        p = doc.add_paragraph()
        if source_section == "Sources consulted" and "http" in line:
            before, url = line.rsplit("http", 1)
            add_inline(p, before, 8.7, MUTED)
            add_hyperlink(p, "http" + url, "http" + url)
        else:
            add_inline(p, line, 10.0)

    doc.core_properties.title = "NexaERP Morocco Market and Production Readiness Report"
    doc.core_properties.subject = "Morocco SME ERP market strategy and production readiness"
    doc.core_properties.author = "NexaERP"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
